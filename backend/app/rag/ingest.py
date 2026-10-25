"""Загрузка документов в базу знаний: любой формат → Markdown (docling, с OCR сканов) → чанки по структуре → эмбеддинги.

HR загружает PDF/DOCX/PPTX через админку — без программиста (вторичная проблема кейса).
"""

import hashlib
import io
import logging
import re
from dataclasses import dataclass
from pathlib import Path

import httpx
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.models import Chunk, Document

log = logging.getLogger(__name__)

HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
EMBED_BATCH = 32  # max_client_batch_size TEI по умолчанию


@dataclass
class Section:
    heading: str
    text: str


# --- конвертация -------------------------------------------------------------------------------------


async def docling_markdown(filename: str, data: bytes) -> str:
    """docling-serve: layout-анализ + OCR (русский) → Markdown со структурой заголовков."""
    async with httpx.AsyncClient(timeout=900) as http:
        r = await http.post(
            f"{settings().docling_url}/v1/convert/file",
            files={"files": (filename, data)},
            data={
                "to_formats": "md",
                "do_ocr": "true",
                "ocr_preset": settings().ocr_preset,
                "ocr_lang": settings().ocr_lang,
                "image_export_mode": "placeholder",
            },
        )
        r.raise_for_status()
        doc = r.json()["document"]
    return doc.get("md_content") or ""


DGM_REL = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds"
R_DM = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}dm"
A_T = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"


def smartart_text(slide, shape) -> str:
    """Текст SmartArt лежит в отдельной части ppt/diagrams/data*.xml — python-pptx его не отдаёт."""
    out = []
    for rel in shape._element.iter(DGM_REL):
        part = slide.part.related_part(rel.get(R_DM))
        root = etree.fromstring(part.blob)
        out.extend(t.text.strip() for t in root.iter(A_T) if t.text and t.text.strip())
    return "\n".join(f"- {t}" for t in out)


async def pptx_markdown(data: bytes) -> str:
    """Слайд = раздел. Учитываются группы, таблицы, SmartArt и текст на картинках (OCR) —
    в корпоративных презентациях содержание часто именно там."""
    prs = Presentation(io.BytesIO(data))
    parts = []
    for n, slide in enumerate(prs.slides, 1):
        texts: list[str] = []

        async def walk(shapes, n=n, slide=slide, texts=texts):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    await walk(shape.shapes)
                    continue
                if shape.has_text_frame and shape.text_frame.text.strip():
                    texts.append(shape.text_frame.text.strip())
                if getattr(shape, "has_table", False) and shape.has_table:
                    texts.append("\n".join(" | ".join(c.text for c in row.cells) for row in shape.table.rows))
                if diagram := smartart_text(slide, shape):
                    texts.append(diagram)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        ocr = await docling_markdown(f"slide{n}.{shape.image.ext}", shape.image.blob)
                        ocr = re.sub(r"<!--.*?-->", "", ocr).strip()
                        if ocr:
                            texts.append(ocr)
                    except Exception:
                        log.warning("OCR картинки на слайде %s не удался", n)

        await walk(slide.shapes)
        title_shape = slide.shapes.title
        title = (
            title_shape.text_frame.text.strip()
            if title_shape is not None and title_shape.has_text_frame
            else ""
        )
        title = " ".join((title or (texts[0] if texts else "")).split())[:120]
        parts.append(f"## Слайд {n}. {title}\n\n" + "\n\n".join(texts))
    return "\n\n".join(parts)


async def to_markdown(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in (".md", ".txt"):
        return data.decode("utf-8")
    if suffix == ".pptx":
        return await pptx_markdown(data)
    return await docling_markdown(filename, data)


# --- чанкинг -----------------------------------------------------------------------------------------


NUMBERED = re.compile(r"^(\d+(?:[.,]\d+)*)[.,]?\s")


def heading_level(md_level: int, title: str) -> int:
    """docling размечает все заголовки сканов одним уровнем, поэтому вложенность берём из нумерации:
    «6.» → 1, «6.4.» → 2, «6.4.1» → 3 (ненумерованные — по уровню Markdown)."""
    m = NUMBERED.match(title)
    return len(re.split(r"[.,]", m.group(1))) if m else md_level


def sections(md: str) -> list[Section]:
    """Разбивка по заголовкам с путём «Раздел › Подраздел» — он идёт в контекст чанка."""
    stack: list[tuple[int, str]] = []
    out: list[Section] = []
    buf: list[str] = []

    def flush():
        text = "\n".join(buf).strip()
        if text:
            out.append(Section(" › ".join(t for _, t in stack), text))
        buf.clear()

    for line in md.splitlines():
        m = HEADING.match(line)
        if m:
            flush()
            title = m.group(2).strip()
            level = heading_level(len(m.group(1)), title)
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, title))
        else:
            buf.append(line)
    flush()
    return out


def split(text: str, max_chars: int) -> list[str]:
    """Пакуем абзацы до max_chars; слишком длинный абзац режем по предложениям."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    pieces: list[str] = []
    for p in paras:
        if len(p) <= max_chars:
            pieces.append(p)
            continue
        cur = ""
        for sent in re.split(r"(?<=[.!?;])\s+", p):
            if cur and len(cur) + len(sent) + 1 > max_chars:
                pieces.append(cur)
                cur = ""
            cur = f"{cur} {sent}".strip()
        if cur:
            pieces.append(cur)
    chunks: list[str] = []
    cur = ""
    for p in pieces:
        if cur and len(cur) + len(p) + 2 > max_chars:
            chunks.append(cur)
            cur = ""
        cur = f"{cur}\n\n{p}".strip()
    if cur:
        chunks.append(cur)
    return chunks


SEPARATOR = re.compile(r"^\|[\s|:-]+\|$")
# Штамп листа в локальных нормативных актах: «Лист 3 Листов 23», «№ измен/дата»
STAMP = re.compile(r"(Лист\w*\s*\d+|Листов\s*\d+|№\s*измен)", re.IGNORECASE)


def _letters(line: str) -> str:
    return re.sub(r"[^a-zа-яё]", "", line.lower())


def strip_running_headers(md: str) -> str:
    """Убирает колонтитулы сканов: они повторяются на каждой странице, с шумом OCR, и засоряют чанки,
    а их шапка иногда распознаётся как заголовок раздела и ломает путь «Раздел › Подраздел»."""
    lines = md.splitlines()
    freq: dict[str, int] = {}
    for ln in lines:
        key = _letters(ln)
        if key:
            freq[key] = freq.get(key, 0) + 1
    running = {k for k, n in freq.items() if n >= 3}
    out: list[str] = []
    for ln in lines:
        key = _letters(ln)
        stripped = ln.strip()
        if STAMP.search(ln):
            continue
        # строка таблицы или «заголовок», повторившиеся на 3+ страницах, — колонтитул (в т.ч. с хвостом шума OCR).
        # Обычный текст не трогаем: «- выходные дни;» может законно встречаться в разных пунктах.
        if (
            key
            and stripped.startswith(("#", "|"))
            and (
                key in running
                or stripped.startswith("#")
                and any(key.startswith(r) for r in running if len(r) > 8)
            )
        ):
            continue
        cells = (
            {c.strip() for c in stripped.strip("|").split("|") if c.strip()}
            if stripped.startswith("|")
            else set()
        )
        if len(cells) == 1 and stripped.count("|") > 2:
            continue  # строка таблицы из одинаковых ячеек — артефакт объединённой шапки штампа
        if SEPARATOR.match(ln.strip()) and not (out and out[-1].strip().startswith("|")):
            continue  # разделитель таблицы, чья шапка удалена выше
        out.append(ln)
    return "\n".join(out)


def chunk_markdown(md: str, title: str, max_chars: int) -> list[tuple[str, str]]:
    md = strip_running_headers(re.sub(r"<!--.*?-->", "", md))
    out = []
    for sec in sections(md):
        heading = f"{title} › {sec.heading}" if sec.heading else title
        for piece in split(sec.text, max_chars):
            out.append((heading, piece))
    return out


# --- эмбеддинги --------------------------------------------------------------------------------------


async def embed(texts: list[str]) -> list[list[float]]:
    """TEI, OpenAI-совместимый /v1/embeddings — тот же код работает с любым провайдером эмбеддингов."""
    out: list[list[float]] = []
    async with httpx.AsyncClient(timeout=300) as http:
        for i in range(0, len(texts), EMBED_BATCH):
            r = await http.post(
                f"{settings().embed_url}/v1/embeddings",
                json={"input": texts[i : i + EMBED_BATCH], "model": "embed"},
            )
            r.raise_for_status()
            out.extend(d["embedding"] for d in sorted(r.json()["data"], key=lambda d: d["index"]))
    return out


def chunk_for_embedding(heading: str, text: str) -> str:
    # Контекст раздела в тексте чанка: дешёвый вариант Contextual Retrieval (Anthropic, 2024) —
    # чанк «4.2 ... не позднее 9:00» без заголовка не находится по запросу «во сколько начинается рабочий день»
    return f"{heading}\n{text}"


# --- пайплайн ----------------------------------------------------------------------------------------


async def ingest(
    s: AsyncSession, filename: str, data: bytes, title: str, site_codes: list[str], max_chars: int
) -> Document:
    sha = hashlib.sha256(data).hexdigest()
    doc = await s.scalar(select(Document).where(Document.sha256 == sha))
    if doc is None:
        doc = Document(title=title, filename=filename, sha256=sha, site_codes=site_codes)
        s.add(doc)
    doc.title, doc.site_codes, doc.status, doc.error = title, site_codes, "processing", None
    await s.commit()
    try:
        # распознанный текст кэшируется по хэшу файла: переиндексация (другой размер чанка) не гоняет OCR заново
        cache = settings().data_dir / "markdown" / f"{sha}.md"
        if cache.exists():
            md = cache.read_text(encoding="utf-8")
        else:
            md = await to_markdown(filename, data)
            cache.parent.mkdir(parents=True, exist_ok=True)
            cache.write_text(md, encoding="utf-8")
        pieces = chunk_markdown(md, title, max_chars)
        vectors = await embed([chunk_for_embedding(h, t) for h, t in pieces])
        await s.execute(delete(Chunk).where(Chunk.document_id == doc.id))
        for i, ((h, t), v) in enumerate(zip(pieces, vectors, strict=True)):
            s.add(Chunk(document_id=doc.id, ord=i, heading=h, text=t, embedding=v))
        doc.status, doc.n_chunks = "ready", len(pieces)
    except Exception as e:
        log.exception("ingest failed: %s", filename)
        doc.status, doc.error = "error", str(e)[:1000]
    await s.commit()
    from app.rag.answer import forget_answers  # answer → search → ingest: импорт здесь, иначе цикл

    forget_answers()
    return doc
